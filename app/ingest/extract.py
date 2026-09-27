"""Text extraction for PDF, DOCX, XLSX/CSV, TXT/MD, HTML and images.

PDFs use the embedded text layer first and fall back to OCR (Hindi + English)
page by page, so scanned government orders are handled too.
"""

import csv
import logging
import re
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

from app.config import settings
from app.lang import normalize

log = logging.getLogger(__name__)

MIN_PAGE_CHARS = 50
WINDOWS_TESSERACT = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")

SUPPORTED_EXTENSIONS = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".csv": "text/csv",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".html": "text/html",
    ".htm": "text/html",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
}


class ExtractionError(Exception):
    pass


@dataclass
class Page:
    number: int | None
    text: str


@dataclass
class Extraction:
    pages: list[Page] = field(default_factory=list)
    page_count: int = 0
    ocr_pages: int = 0
    title: str = ""


def mime_for(filename: str) -> str:
    return SUPPORTED_EXTENSIONS.get(Path(filename).suffix.lower(), "application/octet-stream")


def extract(path: Path) -> Extraction:
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ExtractionError(f"Unsupported file type: {ext}")
    handler = {
        ".pdf": _extract_pdf,
        ".docx": _extract_docx,
        ".xlsx": _extract_xlsx,
        ".csv": _extract_csv,
        ".txt": _extract_text,
        ".md": _extract_text,
        ".html": _extract_html_file,
        ".htm": _extract_html_file,
        ".png": _extract_image,
        ".jpg": _extract_image,
        ".jpeg": _extract_image,
    }[ext]
    result = handler(path)
    result.pages = [Page(p.number, normalize(p.text)) for p in result.pages if p.text and p.text.strip()]
    if not result.pages:
        raise ExtractionError(
            "No text could be extracted. If this is a scanned document, make sure Tesseract OCR "
            "with Hindi ('hin') and English ('eng') language data is installed."
        )
    return result


# ---------------------------------------------------------------- OCR

@lru_cache(maxsize=1)
def _ocr_langs() -> str | None:
    """Configured OCR languages that are actually installed, or None if Tesseract is missing."""
    try:
        import pytesseract

        if settings.tesseract_cmd:
            pytesseract.pytesseract.tesseract_cmd = settings.tesseract_cmd
        elif WINDOWS_TESSERACT.exists():
            pytesseract.pytesseract.tesseract_cmd = str(WINDOWS_TESSERACT)
        available = set(pytesseract.get_languages(config=""))
    except Exception as exc:  # tesseract binary not found
        log.warning("Tesseract OCR not available: %s", exc)
        return None
    wanted = [lang for lang in settings.ocr_langs.split("+") if lang]
    usable = [lang for lang in wanted if lang in available]
    missing = [lang for lang in wanted if lang not in available]
    if missing:
        log.warning("Tesseract language data missing for %s; OCR will use %s", missing, usable or "none")
    return "+".join(usable) or None


def _ocr_image(image) -> str:
    langs = _ocr_langs()
    if not langs:
        return ""
    import pytesseract

    return pytesseract.image_to_string(image, lang=langs)


# ---------------------------------------------------------------- handlers

def _extract_pdf(path: Path) -> Extraction:
    import pymupdf
    from PIL import Image

    try:
        pdf = pymupdf.open(path)
    except Exception as exc:
        raise ExtractionError(f"Could not open PDF (file may be corrupted): {exc}") from exc
    if len(pdf) == 0:
        raise ExtractionError("PDF has no pages (file may be corrupted or truncated).")

    result = Extraction(page_count=len(pdf), title=(pdf.metadata or {}).get("title", "") or "")
    with pdf:
        for index, page in enumerate(pdf):
            text = page.get_text("text")
            if len(text.strip()) < MIN_PAGE_CHARS:
                pix = page.get_pixmap(matrix=pymupdf.Matrix(3, 3), alpha=False)
                image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                ocr_text = _ocr_image(image)
                if len(ocr_text.strip()) > len(text.strip()):
                    text = ocr_text
                    result.ocr_pages += 1
            result.pages.append(Page(index + 1, _clean_pdf_text(text)))
    return result


def _clean_pdf_text(text: str) -> str:
    # Join hyphenated line breaks, then turn single newlines into spaces but keep paragraph breaks.
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    text = re.sub(r"[ \t]*\n[ \t]*\n\s*", "\n\n", text)
    text = re.sub(r"(?<!\n)\n(?!\n)", " ", text)
    text = re.sub(r"\bPage\s*\|\s*\d+\b|\b\d+\s*\|\s*P a g e\b", " ", text)
    return text


def _extract_docx(path: Path) -> Extraction:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))
    title = document.core_properties.title or ""
    return Extraction(pages=[Page(None, "\n\n".join(parts))], page_count=1, title=title)


def _rows_to_text(rows, header=None) -> str:
    lines = []
    for row in rows:
        values = ["" if v is None else str(v).strip() for v in row]
        if not any(values):
            continue
        if header:
            pairs = [f"{h}: {v}" for h, v in zip(header, values) if v]
            lines.append("; ".join(pairs) + ".")
        else:
            lines.append(" | ".join(values))
    return "\n\n".join(lines)


def _extract_xlsx(path: Path) -> Extraction:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    pages = []
    for number, sheet in enumerate(workbook.worksheets, start=1):
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(h).strip() if h is not None else "" for h in rows[0]]
        body = _rows_to_text(rows[1:], header if any(header) else None)
        pages.append(Page(number, f"{sheet.title}\n\n{body}"))
    workbook.close()
    return Extraction(pages=pages, page_count=len(pages))


def _extract_csv(path: Path) -> Extraction:
    with open(path, encoding="utf-8-sig", errors="replace", newline="") as handle:
        rows = list(csv.reader(handle))
    if not rows:
        return Extraction()
    return Extraction(pages=[Page(None, _rows_to_text(rows[1:], rows[0]))], page_count=1)


def _extract_text(path: Path) -> Extraction:
    return Extraction(pages=[Page(None, path.read_text(encoding="utf-8", errors="replace"))], page_count=1)


def html_to_text(html: str) -> tuple[str, str]:
    """Main visible text and title of an HTML page (navigation, headers and footers removed)."""
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "lxml")
    title = ""
    if soup.title and soup.title.string:
        title = soup.title.string.split("|")[0].strip()
    for tag in soup(["script", "style", "nav", "header", "footer", "noscript", "form", "iframe", "svg"]):
        tag.decompose()
    for selector in [".breadcrumb", ".share", ".social", "#skip", ".skip"]:
        for tag in soup.select(selector):
            tag.decompose()
    main = soup.find("main") or soup.find(id="SkipContent") or soup.find("article") or soup.body or soup
    heading = main.find("h1")
    if heading and heading.get_text(strip=True):
        title = heading.get_text(" ", strip=True)
    blocks = []
    for element in main.find_all(["h1", "h2", "h3", "h4", "p", "li", "tr", "td", "div"]):
        if element.name == "div" and element.find(["p", "div", "li", "table"]):
            continue
        if element.name == "td" and element.find_parent("tr"):
            continue
        if element.name == "tr":
            text = " | ".join(c.get_text(" ", strip=True) for c in element.find_all(["td", "th"]))
        else:
            text = element.get_text(" ", strip=True)
        if text and text not in blocks[-1:]:
            blocks.append(text)
    return "\n\n".join(blocks), title


def _extract_html_file(path: Path) -> Extraction:
    text, title = html_to_text(path.read_text(encoding="utf-8", errors="replace"))
    return Extraction(pages=[Page(None, text)], page_count=1, title=title)


def _extract_image(path: Path) -> Extraction:
    from PIL import Image

    with Image.open(path) as image:
        text = _ocr_image(image.convert("RGB"))
    return Extraction(pages=[Page(1, text)], page_count=1, ocr_pages=1)
